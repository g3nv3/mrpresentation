from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Pt


BASE = Path(r"C:\Users\User\Downloads\Telegram Desktop\Литвинов Д.Е. на авторские")

KEEP = {
    "Подсистема сетевой трансляции геокоординат Android-устройства в MR-шлем Pico": {
        "PhoneGpsLocationProvider.cs",
        "PhoneLocationSample.cs",
        "PhoneLocationUdpBroadcaster.cs",
        "PhoneLocationUdpReceiver.cs",
        "Подсистема сетевой трансляции геокоординат Android-устройства в MR-шлем Pico.docx",
    },
    "Подсистема службы переднего плана Android для непрерывного геопозиционирования": {
        "AndroidManifest.xml",
        "AndroidPhoneLocationForegroundService.cs",
        "PhoneLocationForegroundService.java",
        "PhoneLocationForegroundServiceBridge.java",
        "Подсистема службы переднего плана Android для непрерывного геопозиционирования.docx",
    },
    "Фреймворк декларативного управления состояниями пользовательского интерфейса Unity": {
        "IUiControlStateSource.cs",
        "UiControlState.cs",
        "UiControlStateEmitter.cs",
        "UiStateBinder.cs",
        "UiStateRenderer.cs",
        "UiStateTextRenderer.cs",
        "UiStateGraphicColorRenderer.cs",
        "Фреймворк декларативного управления состояниями пользовательского интерфейса Unity.docx",
    },
    "Подсистема видеозахвата и предпросмотра текстуры рендеринга камеры Pico": {
        "PicoCameraFrame.cs",
        "PicoCameraRenderTextureSource.cs",
        "PicoCameraTextureRenderer.cs",
        "Подсистема видеозахвата и предпросмотра текстуры рендеринга камеры Pico.docx",
    },
}


def replace_volume(docx_path: Path, size_bytes: int) -> None:
    document = Document(str(docx_path))
    for paragraph in document.paragraphs:
        if paragraph.text.strip().startswith("Объём программы:"):
            paragraph.clear()
            run = paragraph.add_run(f"Объём программы: {size_bytes} Байт")
            run.font.name = "Times New Roman"
            run.font.size = Pt(14)
            break
    document.save(str(docx_path))


def main() -> None:
    resolved_base = BASE.resolve()
    for dirname, keep_names in KEEP.items():
        directory = BASE / dirname
        if not directory.exists():
            continue

        resolved_directory = directory.resolve()
        if resolved_base not in (resolved_directory, *resolved_directory.parents):
            raise RuntimeError(f"Refusing to edit outside base: {directory}")

        for path in directory.iterdir():
            if path.is_file() and path.name not in keep_names:
                path.unlink()

        source_size = sum(
            path.stat().st_size
            for path in directory.iterdir()
            if path.is_file() and path.suffix.lower() != ".docx"
        )

        for docx_path in directory.glob("*.docx"):
            replace_volume(docx_path, source_size)

        print(f"{directory}: {source_size} bytes")


if __name__ == "__main__":
    main()
