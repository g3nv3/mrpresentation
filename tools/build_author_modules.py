from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


ROOT = Path(r"A:\Games\mrpresentation")
PICO_PROVIDER = Path(r"A:\Games\pico-phone-provider")
EXAMPLES = Path(r"C:\Users\User\Downloads\Telegram Desktop\Литвинов Д.Е. на авторские")
TEMPLATE = EXAMPLES / "QR" / "MR детектор Qr-кодов.docx"
OUT = ROOT / "author_modules"

RIGHT_HOLDER = (
    "Правообладатель: Федеральное государственное казенное образовательное учреждение "
    "высшего образования «Академия Федеральной службы безопасности Российской Федерации»"
)


def source_size(paths: list[Path]) -> int:
    return sum(path.stat().st_size for path in paths)


def clear_body(document: Document) -> None:
    body = document._body._element
    for child in list(body):
        if child.tag.endswith("sectPr"):
            continue
        body.remove(child)


def apply_normal(paragraph, *, bold: bool = False, align=None, after: float | None = None) -> None:
    paragraph.style = "Normal"
    if align is not None:
        paragraph.alignment = align
    if after is not None:
        paragraph.paragraph_format.space_after = Pt(after)
    for run in paragraph.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(14)
        run.font.bold = bold


def add_paragraph(document: Document, text: str = "", *, bold: bool = False, align=None, after=None):
    paragraph = document.add_paragraph()
    if text:
        paragraph.add_run(text)
    apply_normal(paragraph, bold=bold, align=align, after=after)
    return paragraph


def build_docx(output_path: Path, program: str, annotation: list[str], language: str, os_name: str, size_bytes: int) -> None:
    document = Document(str(TEMPLATE))
    clear_body(document)

    add_paragraph(document, "РЕФЕРАТ", align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
    add_paragraph(document)
    add_paragraph(document, RIGHT_HOLDER)
    add_paragraph(document)
    add_paragraph(document, f"Программа: «{program}»")
    add_paragraph(document)
    add_paragraph(document, "Аннотация:", bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, after=0)
    for paragraph_text in annotation:
        add_paragraph(document, paragraph_text)
    add_paragraph(document)
    add_paragraph(document, "Тип ЭВМ: IBM PC", align=WD_ALIGN_PARAGRAPH.LEFT, after=0)
    add_paragraph(document, f"Язык: {language}")
    add_paragraph(document, f"ОС: {os_name}")
    add_paragraph(document, f"Объём программы: {size_bytes} Байт", align=WD_ALIGN_PARAGRAPH.LEFT, after=0)
    add_paragraph(document, "(исходного текста)", align=WD_ALIGN_PARAGRAPH.LEFT, after=0)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_path))


def copy_sources(target_dir: Path, paths: list[Path]) -> None:
    target_dir.mkdir(parents=True, exist_ok=True)
    for path in paths:
        shutil.copy2(path, target_dir / path.name)


def main() -> None:
    if OUT.exists():
        shutil.rmtree(OUT)
    OUT.mkdir(parents=True)

    gps_sources = [
        PICO_PROVIDER / "Assets" / "LocationBridge" / "PhoneGpsLocationProvider.cs",
        PICO_PROVIDER / "Assets" / "LocationBridge" / "PhoneLocationUdpBroadcaster.cs",
        PICO_PROVIDER / "Assets" / "LocationBridge" / "PhoneLocationSample.cs",
        PICO_PROVIDER / "Assets" / "LocationBridge" / "PhoneLocationNetworkSettings.cs",
        PICO_PROVIDER / "Assets" / "LocationBridge" / "LocalNetworkAddresses.cs",
        PICO_PROVIDER / "Assets" / "LocationBridge" / "NetworkAddressDropdown.cs",
        PICO_PROVIDER / "Assets" / "LocationBridge" / "PhoneLocationDebugLogger.cs",
        ROOT / "Assets" / "Project" / "Scripts" / "LocationBridge" / "PhoneLocationUdpReceiver.cs",
        ROOT / "Assets" / "Project" / "Scripts" / "LocationBridge" / "PhoneLocationYandexNavigatorBridge.cs",
        ROOT / "Assets" / "Project" / "Scripts" / "LocationBridge" / "PhoneLocationSample.cs",
    ]

    service_sources = [
        PICO_PROVIDER / "Assets" / "LocationBridge" / "AndroidPhoneLocationForegroundService.cs",
        PICO_PROVIDER
        / "Assets"
        / "Plugins"
        / "Android"
        / "PhoneLocationForegroundService.androidlib"
        / "src"
        / "main"
        / "java"
        / "com"
        / "phonepicoprovider"
        / "locationbridge"
        / "PhoneLocationForegroundService.java",
        PICO_PROVIDER
        / "Assets"
        / "Plugins"
        / "Android"
        / "PhoneLocationForegroundService.androidlib"
        / "src"
        / "main"
        / "java"
        / "com"
        / "phonepicoprovider"
        / "locationbridge"
        / "PhoneLocationForegroundServiceBridge.java",
        PICO_PROVIDER
        / "Assets"
        / "Plugins"
        / "Android"
        / "PhoneLocationForegroundService.androidlib"
        / "AndroidManifest.xml",
        PICO_PROVIDER
        / "Assets"
        / "Plugins"
        / "Android"
        / "PhoneLocationForegroundService.androidlib"
        / "build.gradle",
    ]

    ui_sources = sorted((PICO_PROVIDER / "Assets" / "UiState").glob("*.cs"))

    camera_sources = [
        ROOT / "Assets" / "Project" / "Scripts" / "Camera" / "PicoCameraFrame.cs",
        ROOT / "Assets" / "Project" / "Scripts" / "Camera" / "PicoCameraRenderTextureSource.cs",
        ROOT / "Assets" / "Project" / "Scripts" / "Camera" / "PicoCameraTextureRenderer.cs",
    ]

    gps_dir = OUT / "Подсистема сетевой трансляции геокоординат Android-устройства в MR-шлем Pico"
    ui_dir = OUT / "Фреймворк декларативного управления состояниями пользовательского интерфейса Unity"
    camera_dir = OUT / "Подсистема видеозахвата и предпросмотра текстуры рендеринга камеры Pico"

    build_docx(
        gps_dir / "Подсистема сетевой трансляции геокоординат Android-устройства в MR-шлем Pico.docx",
        "Программное средство сетевой трансляции географических координат Android-устройства в MR-шлем Pico",
        [
            "Программное средство предназначено для получения географических координат Android-устройства и передачи этих данных в приложение смешанной реальности, выполняемое на шлеме Pico. На вход программного средства поступают данные геолокации Android, включающие широту, долготу, высоту, точность, скорость, курс, временную метку и порядковый номер измерения.",
            "Программное средство формирует структурированный пакет координат, проверяет корректность значений и передает данные по локальной сети через UDP. На стороне шлема выполняется прием пакетов, фильтрация по идентификатору устройства и обновление текущей координаты навигационной подсистемы, что позволяет использовать смартфон как внешний источник геопозиционирования для MR-сцены.",
        ],
        "C#",
        "Android, Pico OS",
        source_size(gps_sources),
    )

    build_docx(
        gps_dir / "Подсистема службы переднего плана Android для непрерывного геопозиционирования.docx",
        "Программный модуль службы переднего плана Android для непрерывного геопозиционирования Unity-приложения",
        [
            "Программный модуль предназначен для обеспечения непрерывного получения геолокации на Android-устройстве при работе Unity-приложения. Модуль использует службу переднего плана Android, что позволяет поддерживать процесс геолокации в активном состоянии и корректно взаимодействовать с системными ограничениями фоновой работы.",
            "Модуль включает Java-службу, связующий слой для вызова из Unity и конфигурацию Android-библиотеки. Через программный интерфейс Unity может запускать и останавливать сервис, передавать параметры уведомления, контролировать жизненный цикл геолокационного процесса и использовать полученные координаты в сетевой подсистеме передачи данных на MR-шлем Pico.",
        ],
        "C#, Java",
        "Android",
        source_size(service_sources),
    )
    copy_sources(gps_dir, gps_sources + service_sources)

    build_docx(
        ui_dir / "Фреймворк декларативного управления состояниями пользовательского интерфейса Unity.docx",
        "Программный фреймворк декларативного управления состояниями пользовательского интерфейса Unity",
        [
            "Программный фреймворк предназначен для централизованного управления визуальными состояниями элементов пользовательского интерфейса Unity. На вход фреймворка поступает состояние функционального компонента, после чего связующий слой передает это состояние набору специализированных визуальных обработчиков.",
            "Фреймворк отделяет источник состояния от его отображения и позволяет единообразно изменять текст, цвет, прозрачность, доступность, видимость объектов, спрайты и параметры Animator. Такой подход исключает дублирование логики в отдельных кнопках и панелях, упрощает сопровождение интерфейса и обеспечивает согласованное отображение состояний ожидания, ошибки, активности и недоступности.",
        ],
        "C#",
        "Windows 10/11, Android, Pico OS",
        source_size(ui_sources),
    )
    copy_sources(ui_dir, ui_sources)

    build_docx(
        camera_dir / "Подсистема видеозахвата и предпросмотра текстуры рендеринга камеры Pico.docx",
        "Программный модуль видеозахвата камеры Pico с выводом кадров в текстуру рендеринга и предпросмотром",
        [
            "Программный модуль предназначен для получения кадров с камеры шлема Pico и передачи изображения во внутренний графический конвейер Unity. Модуль запрашивает разрешение на использование камеры, выбирает поддерживаемую конфигурацию устройства, создает сеанс захвата и получает последовательность кадров в формате, пригодном для дальнейшей обработки.",
            "Полученные кадры преобразуются в RenderTexture и могут одновременно выводиться в элемент предпросмотра пользовательского интерфейса. Модуль поддерживает автоматическую инициализацию, корректное освобождение ресурсов, диагностическое логирование и подключение дополнительной обработки изображения, включая распознавание QR-кодов на основе поступающего видеопотока.",
        ],
        "C#",
        "Pico OS",
        source_size(camera_sources),
    )
    copy_sources(camera_dir, camera_sources)

    print(OUT)


if __name__ == "__main__":
    main()
